#!/usr/bin/env python3
"""Read-only audit of candidate workbooks for the QF empirical pilot."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


DATE_TOKENS = ("date", "fecha", "time")


def _serialisable(value):
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, (float, np.floating)):
        return None if not math.isfinite(float(value)) else float(value)
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (pd.Timestamp, np.datetime64)):
        return pd.Timestamp(value).isoformat()
    return str(value)


def _detect_date_column(frame: pd.DataFrame) -> tuple[str | None, pd.Series | None]:
    candidates = [
        col for col in frame.columns
        if any(token in str(col).strip().lower() for token in DATE_TOKENS)
    ]
    if not candidates:
        candidates = list(frame.columns[:3])
    best_col, best_dates, best_count = None, None, 0
    for col in candidates:
        raw = frame[col]
        numeric = pd.to_numeric(raw, errors="coerce")
        eight_digit_share = float(numeric.between(19000101, 21001231).mean())
        if eight_digit_share >= 0.5:
            parsed = pd.to_datetime(
                numeric.round().astype("Int64").astype("string"),
                format="%Y%m%d",
                errors="coerce",
            )
        else:
            parsed = pd.to_datetime(raw, errors="coerce")
        count = int(parsed.notna().sum())
        if count > best_count and count >= max(3, int(0.5 * len(frame))):
            best_col, best_dates, best_count = str(col), parsed, count
    return best_col, best_dates


def _numeric_diagnostics(frame: pd.DataFrame, date_col: str | None) -> dict:
    numeric = frame.select_dtypes(include=[np.number]).copy()
    if date_col in numeric.columns:
        numeric = numeric.drop(columns=[date_col])
    if numeric.empty:
        return {
            "numeric_columns": 0,
            "complete_numeric_columns": 0,
            "nonpositive_columns": 0,
            "columns_with_abs_log_change_gt_40pct": 0,
            "columns_with_abs_simple_change_gt_40pct": 0,
        }

    complete = int((numeric.notna().sum(axis=0) == len(frame)).sum())
    nonpositive = int(((numeric <= 0) & numeric.notna()).any(axis=0).sum())

    positive = numeric.where(numeric > 0)
    abs_log_change = np.log(positive).diff().abs()
    log_jump_cols = int((abs_log_change.max(axis=0, skipna=True) > 0.40).sum())

    abs_simple_change = numeric.pct_change(fill_method=None).abs()
    simple_jump_cols = int((abs_simple_change.max(axis=0, skipna=True) > 0.40).sum())

    return {
        "numeric_columns": int(numeric.shape[1]),
        "complete_numeric_columns": complete,
        "nonpositive_columns": nonpositive,
        "columns_with_abs_log_change_gt_40pct": log_jump_cols,
        "columns_with_abs_simple_change_gt_40pct": simple_jump_cols,
        "median_missing_fraction": float(numeric.isna().mean().median()),
        "max_missing_fraction": float(numeric.isna().mean().max()),
    }


def audit_sheet(path: Path, sheet: str) -> dict:
    frame = pd.read_excel(path, sheet_name=sheet)
    date_col, dates = _detect_date_column(frame)
    result = {
        "sheet": sheet,
        "rows": int(frame.shape[0]),
        "columns": int(frame.shape[1]),
        "column_sample": [str(c) for c in frame.columns[:12]],
        "date_column": date_col,
        "duplicate_column_labels": int(pd.Index(frame.columns).duplicated().sum()),
        "fully_empty_rows": int(frame.isna().all(axis=1).sum()),
        "fully_empty_columns": int(frame.isna().all(axis=0).sum()),
        "numeric": _numeric_diagnostics(frame, date_col),
    }
    if dates is not None:
        valid = dates.dropna()
        result["date_min"] = _serialisable(valid.min()) if not valid.empty else None
        result["date_max"] = _serialisable(valid.max()) if not valid.empty else None
        result["date_valid_rows"] = int(valid.size)
        result["duplicate_dates"] = int(valid.duplicated().sum())
        result["monotonic_increasing_dates"] = bool(valid.is_monotonic_increasing)
    return result


def audit_workbook(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=False)
    sheet_names = list(wb.sheetnames)
    wb.close()
    return {
        "file": path.name,
        "size_bytes": path.stat().st_size,
        "sheets": [audit_sheet(path, sheet) for sheet in sheet_names],
    }


def render_markdown(audits: list[dict]) -> str:
    lines = [
        "# QF empirical-input audit",
        "",
        "Read-only structural and numerical audit. Large price changes are flags for",
        "corporate-action reconciliation, not evidence that an observation is erroneous.",
        "",
        "| Workbook | Sheet | Rows | Columns | Date range | Numeric | Complete numeric | Nonpositive | |log change| > 40% |",
        "|---|---|---:|---:|---|---:|---:|---:|---:|",
    ]
    for book in audits:
        for sheet in book["sheets"]:
            num = sheet["numeric"]
            start = (sheet.get("date_min") or "")[:10]
            end = (sheet.get("date_max") or "")[:10]
            date_range = f"{start}–{end}" if start or end else "not detected"
            lines.append(
                f"| {book['file']} | {sheet['sheet']} | {sheet['rows']} | "
                f"{sheet['columns']} | {date_range} | {num['numeric_columns']} | "
                f"{num['complete_numeric_columns']} | {num['nonpositive_columns']} | "
                f"{num['columns_with_abs_log_change_gt_40pct']} |"
            )
    lines.extend(["", "## Sheet details", ""])
    for book in audits:
        lines.append(f"### {book['file']}")
        lines.append("")
        for sheet in book["sheets"]:
            cols = ", ".join(f"`{c}`" for c in sheet["column_sample"])
            lines.append(
                f"- **{sheet['sheet']}** — date column: `{sheet.get('date_column')}`; "
                f"duplicate dates: {sheet.get('duplicate_dates', 'n/a')}; "
                f"sample columns: {cols}."
            )
        lines.append("")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--json", type=Path, required=True)
    parser.add_argument("--markdown", type=Path, required=True)
    args = parser.parse_args()

    audits = [audit_workbook(path) for path in args.inputs]
    args.json.parent.mkdir(parents=True, exist_ok=True)
    args.json.write_text(json.dumps(audits, indent=2, default=_serialisable), encoding="utf-8")
    args.markdown.write_text(render_markdown(audits), encoding="utf-8")


if __name__ == "__main__":
    main()
